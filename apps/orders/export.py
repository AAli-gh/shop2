"""Orders as an Excel sheet for accounting: one row per order, amounts in Toman, Jalali dates, a totals row."""
import io

from apps.common.admin_jalali import to_jalali

COLUMNS = [
    ("شماره سفارش", 16), ("تاریخ ثبت", 17), ("وضعیت", 14), ("نام گیرنده", 20), ("موبایل گیرنده", 14),
    ("مشتری", 24), ("استان", 12), ("شهر", 12), ("کالاها", 50), ("جمع کالاها", 14), ("تخفیف", 12),
    ("کد تخفیف", 12), ("هزینه ارسال", 12), ("هزینه نصب", 12), ("از کیف پول", 13), ("مبلغ کل", 14),
    ("کد پرداخت", 18), ("کد رهگیری مرسوله", 18),
]
MONEY_TITLES = {"جمع کالاها", "تخفیف", "هزینه ارسال", "هزینه نصب", "از کیف پول", "مبلغ کل"}


def orders_workbook(queryset) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "سفارش‌ها"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "B2"
    sheet.append([title for title, _ in COLUMNS])
    for index, (title, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0D5C75")
        sheet.column_dimensions[get_column_letter(index)].width = width

    orders = queryset.select_related("user", "shipping_address").prefetch_related("items").order_by("-created_at")
    for order in orders:
        address = getattr(order, "shipping_address", None)
        sheet.append([
            order.order_number,
            to_jalali(order.created_at),
            order.get_status_display(),
            address.full_name if address else "",
            address.phone if address else "",
            str(order.user),
            address.province if address else "",
            address.city if address else "",
            "، ".join(f"{item.product_name} × {item.quantity}" for item in order.items.all()),
            order.items_subtotal,
            order.discount_amount,
            order.coupon_code,
            order.shipping_cost,
            order.install_cost,
            order.wallet_used,
            order.total,
            order.payment_ref,
            order.tracking_code,
        ])

    last = sheet.max_row
    money_columns = [index for index, (title, _) in enumerate(COLUMNS, start=1) if title in MONEY_TITLES]
    if last > 1:
        totals = [""] * len(COLUMNS)
        totals[0] = "جمع"
        for index in money_columns:
            letter = get_column_letter(index)
            totals[index - 1] = f"=SUM({letter}2:{letter}{last})"
        sheet.append(totals)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
    for index in money_columns:
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=index).number_format = "#,##0"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
