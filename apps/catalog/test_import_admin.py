import io
import zipfile
from pathlib import Path

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from PIL import Image

from apps.catalog import importer
from apps.catalog.models import Brand, Category, Product

IMPORT_URL = "/admin/catalog/product/import/"


def _xlsx(rows, name="products.xlsx"):
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


def _png_bytes(color="red"):
    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), color).save(buffer, "PNG")
    return buffer.getvalue()


def _zip(files, name="images.zip"):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for file_name, data in files.items():
            archive.writestr(file_name, data)
    return SimpleUploadedFile(name, buffer.getvalue())


@pytest.fixture(autouse=True)
def _media(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / "media"


@pytest.fixture
def fridge(db):
    category = Category.objects.create(name="یخچال و فریزر", slug="fridge")
    return Product.objects.create(category=category, name="یخچال قدیمی", slug="old-fridge", price=30_000_000,
                                  sale_price=28_000_000, stock=0, warranty="۱۸ ماه", is_featured=True)


HEADER = ["شناسه", "نام محصول", "دسته‌بندی", "برند", "قیمت", "قیمت با تخفیف", "موجودی", "فعال", "تصاویر"]


@pytest.mark.django_db
def test_preview_saves_nothing_and_confirm_imports(admin_client, fridge):
    upload = _xlsx([HEADER, ["", "ماشین لباسشویی نو", "ماشین لباسشویی", "برند تازه", "۲۵٬۰۰۰٬۰۰۰", "", "3", "بله", ""]])

    preview = admin_client.post(IMPORT_URL, {"file": upload})

    body = preview.content.decode()
    assert "پیش‌نمایش" in body and "1 محصول جدید" in body
    assert "ماشین لباسشویی" in body and "برند تازه" in body  # new category/brand are announced
    assert Product.objects.count() == 1 and not Brand.objects.exists()
    token = admin_client.session["product_import"]["token"]
    data_path = Path(admin_client.session["product_import"]["data_path"])
    assert data_path.exists()

    response = admin_client.post(IMPORT_URL, {"step": "confirm", "token": token})

    assert response.status_code == 302
    product = Product.objects.get(name="ماشین لباسشویی نو")
    assert (product.price, product.stock, product.brand.name) == (25_000_000, 3, "برند تازه")
    assert not data_path.exists()
    assert "product_import" not in admin_client.session


@pytest.mark.django_db
def test_update_changes_only_the_filled_cells(fridge):
    upload = _xlsx([["شناسه", "قیمت", "موجودی"], ["old-fridge", "32000000", ""]])

    report = importer.run_import(upload, upload.name, commit=True)

    fridge.refresh_from_db()
    assert report.updated == 1 and report.errors == []
    assert fridge.price == 32_000_000
    assert (fridge.sale_price, fridge.stock, fridge.warranty, fridge.is_featured) == (28_000_000, 0, "۱۸ ماه", True)


@pytest.mark.django_db
def test_sale_price_zero_removes_the_discount(fridge):
    upload = _xlsx([["شناسه", "قیمت با تخفیف"], ["old-fridge", "0"]])
    importer.run_import(upload, upload.name, commit=True)
    fridge.refresh_from_db()
    assert fridge.sale_price is None


@pytest.mark.django_db
def test_bad_rows_are_reported_and_good_rows_still_import(fridge):
    upload = _xlsx([
        HEADER,
        ["", "کالای درست", "یخچال و فریزر", "", "1000", "", "1", "", ""],
        ["", "بدون قیمت", "یخچال و فریزر", "", "", "", "", "", ""],
        ["", "قیمت متنی", "یخچال و فریزر", "", "گران", "", "", "", ""],
        ["", "فعال نامعتبر", "یخچال و فریزر", "", "1000", "", "", "شاید", ""],
        ["old-fridge", "", "", "", "100", "200", "", "", ""],
        ["old-fridge", "", "", "", "500", "", "", "", ""],
    ])

    report = importer.run_import(upload, upload.name, commit=True)

    assert report.created == 1
    messages = {row.row: row.message for row in report.errors}
    assert "قیمت" in messages[3]
    assert "عدد" in messages[4]
    assert "بله" in messages[5]
    assert "کمتر" in messages[6]
    assert "ردیف 6" in messages[7]  # duplicate slug in the same file
    fridge.refresh_from_db()
    assert fridge.price == 30_000_000


@pytest.mark.django_db
def test_images_from_zip_become_main_image_and_gallery(fridge, settings):
    upload = _xlsx([["شناسه", "تصاویر"], ["old-fridge", "front.png | photos/side.png"]])
    images = _zip({"front.png": _png_bytes(), "photos/side.png": _png_bytes("blue")})

    preview = importer.run_import(upload, upload.name, images, commit=False)
    assert preview.updated == 1
    assert not Path(settings.MEDIA_ROOT).exists()  # the preview writes no files

    upload.seek(0)
    images.seek(0)
    importer.run_import(upload, upload.name, images, commit=True)

    fridge.refresh_from_db()
    assert fridge.image.name.startswith("products/front")
    assert [Path(i.image.name).name.split("_")[0].split(".")[0] for i in fridge.images.all()] == ["side"]


@pytest.mark.django_db
def test_missing_or_broken_images_are_row_errors(fridge):
    Product.objects.create(category=fridge.category, name="دوم", slug="second", price=1000)
    upload = _xlsx([["شناسه", "تصاویر"], ["old-fridge", "nope.png"], ["second", "bad.png"]])
    images = _zip({"bad.png": b"not an image"})

    report = importer.run_import(upload, upload.name, images, commit=True)

    assert [row.action for row in report.rows] == ["error", "error"]
    assert "nope.png" in report.rows[0].message
    assert "bad.png" in report.rows[1].message


@pytest.mark.django_db
def test_restock_subscribers_are_notified_after_import(fridge, monkeypatch, django_capture_on_commit_callbacks):
    notified = []
    monkeypatch.setattr("apps.orders.emails.restock_available", lambda product: notified.append(product.pk))
    upload = _xlsx([["شناسه", "موجودی"], ["old-fridge", "4"]])

    with django_capture_on_commit_callbacks(execute=True):
        importer.run_import(upload, upload.name, commit=True)

    assert notified == [fridge.pk]


@pytest.mark.django_db
def test_exported_workbook_imports_back_unchanged(admin_client, fridge):
    response = admin_client.post("/admin/catalog/product/", {
        "action": "export_excel", "_selected_action": [fridge.pk],
    })
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")

    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook["محصولات"]
    assert [c.value for c in sheet[1]] == [c.title for c in importer.COLUMNS]
    assert sheet.sheet_view.rightToLeft
    sheet.cell(row=2, column=5, value=31_000_000)  # قیمت
    edited = io.BytesIO()
    workbook.save(edited)

    report = importer.run_import(SimpleUploadedFile("e.xlsx", edited.getvalue()), "e.xlsx", commit=True)

    fridge.refresh_from_db()
    assert report.errors == [] and report.updated == 1
    assert (fridge.price, fridge.sale_price, fridge.warranty, fridge.is_featured) == (31_000_000, 28_000_000, "۱۸ ماه", True)


@pytest.mark.django_db
def test_template_download_and_changelist_button(admin_client):
    response = admin_client.get("/admin/catalog/product/import/template/")
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["محصولات", "فهرست‌ها", "راهنما"]

    assert IMPORT_URL in admin_client.get("/admin/catalog/product/").content.decode()


@pytest.mark.django_db
def test_unreadable_upload_shows_an_error(admin_client):
    response = admin_client.post(IMPORT_URL, {"file": SimpleUploadedFile("p.xlsx", b"not a workbook")})
    assert "فایل Excel خوانده نشد" in response.content.decode()
    assert "product_import" not in admin_client.session


@pytest.mark.django_db
def test_stale_confirm_token_is_rejected(admin_client):
    response = admin_client.post(IMPORT_URL, {"step": "confirm", "token": "forged"})
    assert response.status_code == 302 and not Product.objects.exists()


@pytest.mark.django_db
def test_import_needs_add_and_change_permission(client, django_user_model):
    from django.contrib.auth.models import Permission

    staff = django_user_model.objects.create_user(email="editor@example.com", password="x", is_staff=True)
    staff.user_permissions.add(Permission.objects.get(codename="view_product"))
    client.force_login(staff)

    assert client.get(IMPORT_URL).status_code == 403
    assert IMPORT_URL not in client.get("/admin/catalog/product/").content.decode()
