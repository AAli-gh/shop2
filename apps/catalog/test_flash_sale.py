import io
import json
import re
from datetime import timedelta

import openpyxl
import pytest
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from PIL import Image

from apps.catalog import importer, services
from apps.catalog.models import Category, Product, ProductVariant
from apps.catalog.video import embed_url


@pytest.fixture
def shop(db):
    category = Category.objects.create(name="یخچال و فریزر", slug="fridge")
    live = Product.objects.create(category=category, name="یخچال تخفیف‌دار", slug="live", price=50_000_000,
                                  sale_price=40_000_000, stock=4, sale_ends_at=timezone.now() + timedelta(hours=5))
    ended = Product.objects.create(category=category, name="یخچال تمام‌شده", slug="ended", price=30_000_000,
                                   sale_price=25_000_000, stock=2, sale_ends_at=timezone.now() - timedelta(hours=1))
    plain = Product.objects.create(category=category, name="یخچال ساده", slug="plain", price=20_000_000, stock=0)
    return live, ended, plain


def test_a_discount_only_counts_inside_its_window(shop):
    live, ended, plain = shop

    assert (live.on_sale, live.flash_sale, live.effective_price, live.discount_percent) == (True, True, 40_000_000, 20)
    assert (ended.on_sale, ended.effective_price, ended.discount_percent) == (False, 30_000_000, 0)

    live.sale_starts_at = timezone.now() + timedelta(hours=1)  # scheduled for later today
    assert live.on_sale is False and live.effective_price == 50_000_000

    plain.sale_price = 18_000_000  # no dates: a plain discount is always on
    assert plain.on_sale is True and plain.flash_sale is False


def test_variants_follow_the_product_window(shop):
    live, ended, _ = shop
    running = ProductVariant.objects.create(product=live, color_name="سفید", price=52_000_000, sale_price=41_000_000, stock=3)
    finished = ProductVariant.objects.create(product=ended, color_name="سفید", price=31_000_000, sale_price=26_000_000, stock=1)

    assert running.effective_price == 41_000_000
    assert finished.effective_price == 31_000_000 and finished.discount_percent == 0


def test_only_live_discounts_are_listed(shop):
    live, ended, _ = shop
    assert [p.slug for p in services.on_sale_products()] == ["live"]
    assert [p.slug for p in services.flash_sale_products()] == ["live"]

    live.sale_ends_at = timezone.now() - timedelta(minutes=1)
    live.save()
    assert services.flash_sale_products() == [] and services.on_sale_products() == []


@pytest.mark.django_db
def test_the_home_page_counts_down_to_the_first_deadline(client, shop):
    live, _, _ = shop

    page = client.get("/").content.decode()

    assert "شگفت‌انگیز امروز" in page
    assert f'data-countdown="{timezone.localtime(live.sale_ends_at).isoformat()}"' in page  # the browser counts down to it
    assert "یخچال تخفیف‌دار" in page


@pytest.mark.django_db
def test_listing_filters_for_stock_and_discounts(client, shop):
    in_stock_only = client.get("/products/", {"stock": "1"}).content.decode()
    assert "یخچال ساده" not in in_stock_only  # out of stock
    assert "یخچال تخفیف‌دار" in in_stock_only

    on_sale_only = client.get("/products/", {"sale": "1"}).content.decode()
    assert "یخچال تخفیف‌دار" in on_sale_only
    assert "یخچال تمام‌شده" not in on_sale_only  # its discount ended
    assert 'name="sale" value="1" checked' in on_sale_only


@pytest.mark.django_db
def test_structured_data_says_when_the_discount_ends(client, shop):
    live, _, _ = shop
    page = client.get(live.get_absolute_url()).content.decode()
    blocks = [json.loads(b) for b in re.findall(r'<script type="application/ld\+json">(.*?)</script>', page, re.S)]
    offers = next(b for b in blocks if b.get("@type") == "Product")["offers"]
    assert offers["priceValidUntil"] == live.sale_ends_at.date().isoformat()


@pytest.mark.parametrize("url, expected", [
    ("https://www.aparat.com/v/AbC123", "https://www.aparat.com/video/video/embed/videohash/AbC123/vt/frame"),
    ("https://www.youtube.com/watch?v=dQw4w9WgXcQ", "https://www.youtube.com/embed/dQw4w9WgXcQ"),
    ("https://youtu.be/dQw4w9WgXcQ", "https://www.youtube.com/embed/dQw4w9WgXcQ"),
    ("https://example.com/video.mp4", ""),
    ("", ""),
])
def test_known_video_hosts_play_in_the_page(url, expected):
    assert embed_url(url) == expected


@pytest.mark.django_db
def test_the_product_page_embeds_the_video_and_opens_photos(client, shop, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    live, _, _ = shop
    live.video_url = "https://www.aparat.com/v/AbC123"
    photo = io.BytesIO()
    Image.new("RGB", (40, 40), "white").save(photo, "PNG")
    live.image.save("fridge.png", ContentFile(photo.getvalue()))
    live.save()

    page = client.get(live.get_absolute_url()).content.decode()

    assert '<iframe src="https://www.aparat.com/video/video/embed/videohash/AbC123/vt/frame"' in page
    assert "data-main-image" in page and "cursor-zoom-in" in page


def test_import_reads_jalali_sale_dates_and_the_cost_price(shop):
    workbook = openpyxl.Workbook()
    workbook.active.append(["شناسه", "شروع تخفیف", "پایان تخفیف", "قیمت خرید"])
    workbook.active.append(["plain", "۱۴۰۵/۰۷/۰۵", "1405/07/09 23:59", "۱۵,۰۰۰,۰۰۰"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    report = importer.run_import(SimpleUploadedFile("p.xlsx", buffer.getvalue()), "p.xlsx", commit=True)

    assert report.errors == []
    product = Product.objects.get(slug="plain")
    starts, ends = timezone.localtime(product.sale_starts_at), timezone.localtime(product.sale_ends_at)
    assert (starts.year, starts.month, starts.day) == (2026, 9, 27)  # 1405/07/05
    assert (ends.hour, ends.minute) == (23, 59)
    assert product.cost_price == 15_000_000


def test_import_reports_a_bad_date(shop):
    workbook = openpyxl.Workbook()
    workbook.active.append(["شناسه", "پایان تخفیف"])
    workbook.active.append(["plain", "دیروز"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    report = importer.run_import(SimpleUploadedFile("p.xlsx", buffer.getvalue()), "p.xlsx", commit=True)

    assert "تاریخ" in report.errors[0].message


def test_export_writes_the_sale_window_back_in_jalali(shop):
    live, _, _ = shop
    workbook = openpyxl.load_workbook(io.BytesIO(importer.build_workbook(Product.objects.filter(pk=live.pk))))
    row = {title: cell.value for (title, _), cell in zip(
        [(c.title, c.width) for c in importer.COLUMNS], workbook["محصولات"][2])}
    assert re.match(r"^1[34]\d\d/\d\d/\d\d \d\d:\d\d$", row["پایان تخفیف"])
    assert not row["شروع تخفیف"]  # never started explicitly, so the cell stays empty
