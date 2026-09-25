"""Bulk product import from Excel (xlsx) or CSV, with an optional ZIP of photos.

Headers may be the Persian ones of the downloadable template or the English field
names. A row whose slug (شناسه — or, when blank, the slug made from the name) matches
an existing product updates it; any other row creates a product.

On update, an empty cell leaves the current value alone and a column missing from the
file is never touched, so a sheet with only «شناسه» and «قیمت» is a bulk price change.
`run_import(..., commit=False)` does every step inside a transaction and rolls it back,
which is the preview shown before anything is saved.

An xlsx file may carry a second sheet, «تنوع‌ها», with one row per colour/option of a product
(matched by the product's slug, then by the variant's SKU or its colour + option). The products
sheet is applied first, so a new product and its variants can arrive in the same file.
"""
import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath

from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone
from django.utils.text import slugify

from apps.catalog.models import Brand, Category, Product, ProductImage, ProductVariant
from apps.common.text import normalize_search

MAX_ROWS = 5000
MAX_IMAGE_BYTES = 5 * 1024 * 1024

FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
SEP_RE = re.compile(r"[,،٬\s]")
TRUE_WORDS = {"1", "true", "yes", "y", "بله", "آری", "فعال", "دارد", "✓"}
FALSE_WORDS = {"0", "false", "no", "n", "خیر", "نه", "غیرفعال", "ندارد"}


@dataclass(frozen=True)
class Column:
    key: str
    title: str  # header in the template
    help: str
    aliases: tuple = ()
    width: int = 16


COLUMNS = [
    Column("slug", "شناسه", "شناسه یکتای محصول در آدرس صفحه. برای ویرایش محصول موجود پر شود؛ خالی = ساخت از روی نام.", width=24),
    Column("name", "نام محصول", "برای محصول جدید الزامی.", width=34),
    Column("category", "دسته‌بندی", "برای محصول جدید الزامی. دسته‌ای که وجود نداشته باشد ساخته می‌شود.", ("دسته",), 20),
    Column("brand", "برند", "برندی که وجود نداشته باشد ساخته می‌شود.", width=16),
    Column("price", "قیمت", "به تومان؛ برای محصول جدید الزامی.", ("قیمت (تومان)",), 14),
    Column("sale_price", "قیمت با تخفیف", "به تومان و کمتر از قیمت. عدد ۰ تخفیف را حذف می‌کند.", ("قیمت تخفیف",), 14),
    Column("sale_starts_at", "شروع تخفیف", "تاریخ شمسی مثل ۱۴۰۵/۰۷/۰۵ یا با ساعت ۱۴۰۵/۰۷/۰۵ ۰۹:۰۰؛ خالی = از همین حالا.",
           ("شروع فروش ویژه",), 18),
    Column("sale_ends_at", "پایان تخفیف",
           "پایان تخفیف؛ با پر کردن آن کالا با شمارش معکوس در «شگفت‌انگیز امروز» می‌آید. خالی = بدون مهلت.",
           ("پایان فروش ویژه",), 18),
    Column("cost_price", "قیمت خرید", "قیمت تمام‌شده به تومان؛ فقط برای گزارش سود و به مشتری نشان داده نمی‌شود.", (), 14),
    Column("stock", "موجودی", "تعداد در انبار.", width=10),
    Column("is_active", "فعال", "بله یا خیر — محصول غیرفعال در سایت دیده نمی‌شود.", ("active",), 8),
    Column("is_featured", "ویژه", "بله یا خیر — نمایش در ویترین صفحه اول.", ("featured", "ویترین"), 8),
    Column("warranty", "گارانتی", "مثلاً ۲۴ ماه گارانتی رسمی.", width=24),
    Column("country", "کشور سازنده", "", ("country of origin",), 14),
    Column("energy_class", "رده انرژی", "مثلاً A+ یا A++.", ("energy class",), 10),
    Column("weight_kg", "وزن (کیلوگرم)", "عدد، مثلاً 62.5", ("وزن", "weight"), 12),
    Column("requires_install", "نیاز به نصب", "بله یا خیر.", ("requires install",), 10),
    Column("installment_months", "ماه‌های اقساط", "۰ = بدون فروش قسطی.", width=12),
    Column("installment_down_payment", "پیش‌پرداخت اقساط", "به تومان.", width=14),
    Column("video_url", "لینک ویدیو", "آدرس کامل ویدیو در آپارات یا یوتیوب.", width=24),
    Column("description", "توضیحات", "", width=40),
    Column("images", "تصاویر", "نام فایل‌ها داخل ZIP تصاویر، جداشده با |. اولی تصویر اصلی، بقیه گالری (جایگزین گالری فعلی).",
           ("image", "تصویر"), 28),
]
COLUMN_BY_KEY = {column.key: column for column in COLUMNS}
BOOL_KEYS = {"is_active", "is_featured", "requires_install"}
INT_KEYS = {"price", "sale_price", "stock", "installment_months", "installment_down_payment", "cost_price"}
DATE_KEYS = {"sale_starts_at", "sale_ends_at"}
TEXT_KEYS = {"name", "warranty", "country", "energy_class", "video_url", "description"}

PRODUCTS_SHEET = "محصولات"
VARIANTS_SHEET = "تنوع‌ها"
VARIANT_COLUMNS = [
    Column("product", "شناسه محصول", "شناسهٔ محصولی که این رنگ یا ظرفیت مال آن است (ستون «شناسه» در برگهٔ محصولات).",
           ("شناسه", "product slug"), 28),
    Column("color_name", "رنگ", "مثلاً سفید یا استیل. «رنگ» یا «گزینه»، حداقل یکی لازم است.", ("color",), 14),
    Column("color_hex", "کد رنگ", "برای نمونهٔ رنگ در سایت، مثلاً #C0C0C0 (اختیاری).", ("color code",), 12),
    Column("option", "گزینه", "ظرفیت، سایز یا مدل؛ مثلاً ۲۴ فوت.", ("ظرفیت", "مدل"), 16),
    Column("sku", "کد کالا", "کد انبار این تنوع (اختیاری). اگر پر باشد، تنوع با همین کد پیدا و ویرایش می‌شود.", (), 16),
    Column("price", "قیمت", "به تومان؛ برای تنوع جدید الزامی.", ("قیمت (تومان)",), 14),
    Column("sale_price", "قیمت با تخفیف", "به تومان و کمتر از قیمت. عدد ۰ تخفیف را حذف می‌کند.", ("قیمت تخفیف",), 14),
    Column("stock", "موجودی", "تعداد در انبار.", width=10),
    Column("is_active", "فعال", "بله یا خیر.", ("active",), 8),
    Column("sort_order", "ترتیب", "ترتیب نمایش در صفحهٔ محصول؛ عدد کوچک‌تر اول.", ("sort order",), 8),
    Column("image", "تصویر", "نام یک فایل داخل ZIP تصاویر، مخصوص همین رنگ (اختیاری).", ("images", "تصاویر"), 22),
]
VARIANT_COLUMN_BY_KEY = {column.key: column for column in VARIANT_COLUMNS}


class ImportFileError(Exception):
    """The file as a whole can't be read (wrong format, no header, too many rows)."""


class RowError(Exception):
    pass


@dataclass
class RowResult:
    row: int
    name: str
    action: str  # "create" | "update" | "error"
    message: str = ""
    sheet: str = PRODUCTS_SHEET


@dataclass
class ImportReport:
    rows: list = field(default_factory=list)
    new_categories: list = field(default_factory=list)
    new_brands: list = field(default_factory=list)
    ignored_columns: list = field(default_factory=list)
    committed: bool = False

    def _count(self, action, sheet=PRODUCTS_SHEET):
        return sum(1 for row in self.rows if row.action == action and row.sheet == sheet)

    @property
    def created(self):
        return self._count("create")

    @property
    def updated(self):
        return self._count("update")

    @property
    def variants_created(self):
        return self._count("create", VARIANTS_SHEET)

    @property
    def variants_updated(self):
        return self._count("update", VARIANTS_SHEET)

    @property
    def saved(self):
        """Rows of both sheets that will be (or were) written."""
        return sum(1 for row in self.rows if row.action != "error")

    @property
    def errors(self):
        return [row for row in self.rows if row.action == "error"]


# ---------------------------------------------------------------- reading

def _header_key(title) -> str:
    return normalize_search(str(title or "").replace("_", " ")).replace(" ", "")


def _lookup(columns) -> dict:
    return {_header_key(name): column.key for column in columns for name in (column.key, column.title, *column.aliases)}


_HEADER_LOOKUP = _lookup(COLUMNS)
_VARIANT_HEADER_LOOKUP = _lookup(VARIANT_COLUMNS)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores 1200000 as 1200000.0
    return str(value).strip()


def _table(rows):
    """(header, data rows) of a sheet, or None when it has no header row."""
    header = None
    data = []
    for cells in rows:
        if header is None:
            if any(cells):
                header = cells
            continue
        if any(cells):
            data.append(cells)
            if len(data) > MAX_ROWS:
                raise ImportFileError(f"حداکثر {MAX_ROWS} ردیف در هر برگه؛ فایل را چند بخش کنید.")
    return (header, data) if header is not None else None


# sheets of the downloadable template that hold no data to import
_REFERENCE_SHEETS = {_header_key("فهرست‌ها"), _header_key("راهنما")}


def read_sheets(file_obj, filename: str):
    """The products table and the variants table ((header, rows) or None) of an xlsx or CSV file."""
    name = (filename or "").lower()
    products = variants = None
    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl

        try:
            workbook = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for a damaged or non-xlsx file
            raise ImportFileError("فایل Excel خوانده نشد؛ آن را با فرمت xlsx ذخیره کنید.") from exc
        for sheet in workbook.worksheets:
            rows = ([_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True))
            title = _header_key(sheet.title)
            if title in (_header_key(VARIANTS_SHEET), "variants"):
                variants = _table(rows)
            elif title not in _REFERENCE_SHEETS and products is None:
                products = _table(rows)
    elif name.endswith(".csv"):
        try:
            text = file_obj.read().decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ImportFileError("فایل CSV باید با کدگذاری UTF-8 ذخیره شده باشد.") from exc
        products = _table([cell.strip() for cell in row] for row in csv.reader(io.StringIO(text)))
    else:
        raise ImportFileError("فقط فایل Excel (xlsx) یا CSV پذیرفته می‌شود.")
    if products is None and variants is None:
        raise ImportFileError("فایل خالی است.")
    return products, variants


# ---------------------------------------------------------------- values

def _to_int(value: str, title: str):
    text = SEP_RE.sub("", value.translate(FA_DIGITS))
    if not text:
        return None
    try:
        number = Decimal(text)
    except InvalidOperation:
        raise RowError(f"«{title}» باید عدد باشد («{value}»).")
    if number != number.to_integral_value() or number < 0:
        raise RowError(f"«{title}» باید عدد صحیح و بدون علامت منفی باشد («{value}»).")
    return int(number)


def _to_datetime(value: str, title: str):
    """Jalali (۱۴۰۵/۰۷/۰۵ ۰۹:۰۰) or Gregorian dates, with or without a time."""
    text = value.translate(FA_DIGITS).strip().replace("-", "/")
    if not text:
        return None
    date_part, _, time_part = text.partition(" ")
    try:
        year, month, day = (int(part) for part in date_part.split("/"))
        clock = [int(part) for part in time_part.split(":")[:2]] if time_part.strip() else []
        hour, minute = (clock + [0, 0])[:2]
        if year < 1700:  # Jalali years are far below Gregorian ones
            import jdatetime

            moment = jdatetime.datetime(year, month, day, hour, minute).togregorian()
        else:
            from datetime import datetime

            moment = datetime(year, month, day, hour, minute)
    except (ValueError, TypeError):
        raise RowError(f"«{title}» باید تاریخ باشد، مثل ۱۴۰۵/۰۷/۰۵ یا ۱۴۰۵/۰۷/۰۵ ۰۹:۰۰ («{value}»).")
    return timezone.make_aware(moment) if timezone.is_naive(moment) else moment


def _to_decimal(value: str, title: str):
    text = value.translate(FA_DIGITS).replace("٫", ".").replace("/", ".").strip()
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        raise RowError(f"«{title}» باید عدد باشد («{value}»).")


def _to_bool(value: str, title: str):
    text = value.translate(FA_DIGITS).strip().lower()
    if not text:
        return None
    if text in TRUE_WORDS:
        return True
    if text in FALSE_WORDS:
        return False
    raise RowError(f"«{title}» باید «بله» یا «خیر» باشد («{value}»).")


def _unique_slug(model, name: str) -> str:
    base = slugify(name, allow_unicode=True) or "item"
    slug, n = base, 2
    while model.objects.filter(slug=slug).exists():
        slug, n = f"{base}-{n}", n + 1
    return slug


class _Lookups:
    """Categories and brands by folded name, created on first use."""

    def __init__(self, report: ImportReport):
        self.report = report
        self.categories = {normalize_search(c.name): c for c in Category.objects.all()}
        self.brands = {normalize_search(b.name): b for b in Brand.objects.all()}

    def category(self, name: str) -> Category:
        key = normalize_search(name)
        if key not in self.categories:
            self.categories[key] = Category.objects.create(name=name, slug=_unique_slug(Category, name))
            self.report.new_categories.append(name)
        return self.categories[key]

    def brand(self, name: str) -> Brand:
        key = normalize_search(name)
        if key not in self.brands:
            self.brands[key] = Brand.objects.create(name=name, slug=_unique_slug(Brand, name))
            self.report.new_brands.append(name)
        return self.brands[key]


class _Images:
    """Photos from the uploaded ZIP, looked up by file name (folders and letter case ignored)."""

    def __init__(self, zip_file):
        self.archive = None
        self.entries = {}
        if zip_file is None:
            return
        try:
            self.archive = zipfile.ZipFile(zip_file)
        except zipfile.BadZipFile as exc:
            raise ImportFileError("فایل تصاویر ZIP معتبر نیست.") from exc
        for info in self.archive.infolist():
            if not info.is_dir():
                self.entries.setdefault(PurePosixPath(info.filename).name.lower(), info)

    def load(self, name: str) -> tuple[str, bytes]:
        from PIL import Image

        if self.archive is None:
            raise RowError("ستون «تصاویر» پر است ولی فایل ZIP تصاویر انتخاب نشده.")
        info = self.entries.get(PurePosixPath(name.replace("\\", "/")).name.lower())
        if info is None:
            raise RowError(f"تصویر «{name}» در فایل ZIP نیست.")
        if info.file_size > MAX_IMAGE_BYTES:
            raise RowError(f"تصویر «{name}» بزرگ‌تر از ۵ مگابایت است.")
        data = self.archive.read(info)
        try:
            with Image.open(io.BytesIO(data)) as image:
                image.verify()
        except Exception:
            raise RowError(f"«{name}» فایل تصویر معتبری نیست.")
        return PurePosixPath(info.filename).name, data


# ---------------------------------------------------------------- import

def run_import(data_file, filename: str, images_zip=None, *, commit: bool) -> ImportReport:
    """Import every valid row; rows with problems are reported and skipped.

    With commit=False nothing is kept (not even photo files), so the report is a preview.
    """
    products_table, variants_table = read_sheets(data_file, filename)
    report = ImportReport()
    positions = _positions(products_table, _HEADER_LOOKUP, report, "")
    variant_positions = _positions(variants_table, _VARIANT_HEADER_LOOKUP, report, f"{VARIANTS_SHEET}: ")
    product_rows = products_table[1] if products_table else []
    variant_rows = variants_table[1] if variants_table else []
    if products_table and "name" not in positions and "slug" not in positions and (product_rows or not variants_table):
        raise ImportFileError("ستون «نام محصول» یا «شناسه» در ردیف اول فایل پیدا نشد؛ از فایل نمونه استفاده کنید.")
    if variant_rows and "product" not in variant_positions:
        raise ImportFileError(f"در برگهٔ «{VARIANTS_SHEET}» ستون «شناسه محصول» پیدا نشد؛ از فایل نمونه استفاده کنید.")

    images = _Images(images_zip)
    lookups = _Lookups(report)
    seen_slugs = {}
    seen_variants = {}
    touched = {}  # product id -> it was out of stock before its variants changed

    with transaction.atomic():
        for offset, cells in enumerate(product_rows):
            row_number = offset + 2  # the header is row 1
            values = {key: (cells[i] if i < len(cells) else "") for key, i in positions.items()}
            savepoint = transaction.savepoint()
            try:
                action, product = _apply_row(values, lookups, images, seen_slugs, row_number, commit=commit)
            except RowError as exc:
                transaction.savepoint_rollback(savepoint)
                report.rows.append(RowResult(row_number, values.get("name") or values.get("slug", ""), "error", str(exc)))
            else:
                transaction.savepoint_commit(savepoint)
                report.rows.append(RowResult(row_number, product.name, action))

        for offset, cells in enumerate(variant_rows):
            row_number = offset + 2
            values = {key: (cells[i] if i < len(cells) else "") for key, i in variant_positions.items()}
            savepoint = transaction.savepoint()
            try:
                action, variant = _apply_variant_row(values, images, seen_variants, touched, row_number, commit=commit)
            except RowError as exc:
                transaction.savepoint_rollback(savepoint)
                label = " · ".join(v for v in (values.get("product"), values.get("color_name"), values.get("option")) if v)
                report.rows.append(RowResult(row_number, label, "error", str(exc), VARIANTS_SHEET))
            else:
                transaction.savepoint_commit(savepoint)
                report.rows.append(RowResult(row_number, f"{variant.product.name} — {variant.label}", action, "", VARIANTS_SHEET))

        _notify_restocked(touched)
        if commit:
            report.committed = True
        else:
            transaction.set_rollback(True)
    return report


def _positions(table, lookup: dict, report: ImportReport, prefix: str) -> dict:
    """Column index of each known field in the header; unknown titles are reported as ignored."""
    positions = {}
    if not table:
        return positions
    for index, title in enumerate(table[0]):
        key = lookup.get(_header_key(title))
        if key is None:
            if title:
                report.ignored_columns.append(prefix + title)
        else:
            positions.setdefault(key, index)
    return positions


def _notify_restocked(touched: dict) -> None:
    """Restock subscribers hear about products that were sold out before their variants came in."""
    from apps.orders.emails import restock_available

    for product_id, was_out_of_stock in touched.items():
        if not was_out_of_stock:
            continue
        product = Product.objects.filter(pk=product_id, stock__gt=0).first()
        if product is not None:
            transaction.on_commit(lambda p=product: restock_available(p))


def _apply_variant_row(values: dict, images: "_Images", seen: dict, touched: dict, row_number: int, *, commit: bool):
    def given(key):
        return values.get(key, "") != ""

    def title(key):
        return VARIANT_COLUMN_BY_KEY[key].title

    slug = values.get("product", "")
    if not slug:
        raise RowError("«شناسه محصول» خالی است.")
    product = Product.objects.filter(slug=slug).first()
    if product is None:
        raise RowError(f"محصولی با شناسهٔ «{slug}» پیدا نشد؛ اول آن را در برگهٔ «{PRODUCTS_SHEET}» بسازید.")

    color, option, sku = values.get("color_name", ""), values.get("option", ""), values.get("sku", "")
    identity = (product.pk, f"sku:{sku}" if sku else f"{color}|{option}")
    if identity in seen:
        raise RowError(f"این تنوع در ردیف {seen[identity]} هم آمده است.")
    seen[identity] = row_number

    variant = product.variants.filter(sku=sku).first() if sku else None
    if variant is None and (color or option):
        variant = product.variants.filter(color_name=color, option=option).first()
    creating = variant is None
    if creating:
        if not (color or option):
            raise RowError("برای تنوع جدید «رنگ» یا «گزینه» لازم است.")
        if not given("price"):
            raise RowError("برای تنوع جدید «قیمت» لازم است.")
        variant = ProductVariant(product=product)
    touched.setdefault(product.pk, product.stock == 0)

    for key in ("color_name", "color_hex", "option", "sku"):
        if given(key):
            setattr(variant, key, values[key])
    for key in ("price", "stock", "sort_order"):
        if given(key):
            setattr(variant, key, _to_int(values[key], title(key)))
    if given("sale_price"):
        variant.sale_price = _to_int(values["sale_price"], title("sale_price")) or None
    if given("is_active"):
        variant.is_active = _to_bool(values["is_active"], title("is_active"))
    photo = images.load(values["image"]) if given("image") else None

    try:
        variant.full_clean(exclude=["image"])
    except ValidationError as exc:
        raise RowError(" ".join(_validation_messages(exc, ProductVariant)))

    if commit and photo:
        variant.image.save(photo[0], ContentFile(photo[1]), save=False)
    variant.save()  # also refreshes the product's price range and total stock
    return ("create" if creating else "update"), variant


def _apply_row(values: dict, lookups: _Lookups, images: _Images, seen_slugs: dict, row_number: int, *, commit: bool):
    name = values.get("name", "")
    slug = values.get("slug", "") or slugify(name, allow_unicode=True)
    if not slug:
        raise RowError("نام محصول یا شناسه لازم است.")
    if slug in seen_slugs:
        raise RowError(f"شناسه «{slug}» در ردیف {seen_slugs[slug]} هم آمده است.")
    seen_slugs[slug] = row_number

    product = Product.objects.filter(slug=slug).first()
    creating = product is None
    if creating:
        missing = [COLUMN_BY_KEY[key].title for key in ("name", "category", "price") if not values.get(key)]
        if missing:
            raise RowError(f"برای محصول جدید این ستون‌ها لازم است: {'، '.join(missing)}.")
        product = Product(slug=slug)
    was_out_of_stock = not creating and product.stock == 0

    def given(key):
        return values.get(key, "") != ""

    if not creating and product.has_variants and any(given(key) for key in ("price", "sale_price", "stock")):
        raise RowError("این محصول تنوع (رنگ یا ظرفیت) دارد؛ قیمت و موجودی را در صفحه محصول، بخش تنوع‌ها تغییر دهید.")

    for key in TEXT_KEYS:
        if given(key):
            setattr(product, key, values[key])
    for key in INT_KEYS - {"sale_price"}:
        if given(key):
            setattr(product, key, _to_int(values[key], COLUMN_BY_KEY[key].title))
    for key in BOOL_KEYS:
        if given(key):
            setattr(product, key, _to_bool(values[key], COLUMN_BY_KEY[key].title))
    if given("sale_price"):
        product.sale_price = _to_int(values["sale_price"], COLUMN_BY_KEY["sale_price"].title) or None
    for key in DATE_KEYS:
        if given(key):
            setattr(product, key, _to_datetime(values[key], COLUMN_BY_KEY[key].title))
    if given("weight_kg"):
        product.weight_kg = _to_decimal(values["weight_kg"], COLUMN_BY_KEY["weight_kg"].title)
    if given("category"):
        product.category = lookups.category(values["category"])
    if given("brand"):
        product.brand = lookups.brand(values["brand"])

    photos = []
    if given("images"):
        names = [part.strip() for part in re.split(r"[|،]", values["images"]) if part.strip()]
        photos = [images.load(photo) for photo in names]

    try:
        product.full_clean(exclude=["image"])
    except ValidationError as exc:
        raise RowError(" ".join(_validation_messages(exc)))

    if commit and photos:
        file_name, data = photos[0]
        product.image.save(file_name, ContentFile(data), save=False)
    product.save()
    if commit and photos:
        product.images.all().delete()
        for order, (file_name, data) in enumerate(photos[1:]):
            gallery = ProductImage(product=product, sort_order=order, alt=product.name)
            gallery.image.save(file_name, ContentFile(data), save=False)
            gallery.save()

    if was_out_of_stock and product.stock > 0:
        from apps.orders.emails import restock_available

        transaction.on_commit(lambda p=product: restock_available(p))
    return ("create" if creating else "update"), product


def _validation_messages(exc: ValidationError, model=Product) -> list[str]:
    if not hasattr(exc, "error_dict"):
        return list(exc.messages)
    messages = []
    for field_name, errors in exc.error_dict.items():
        try:
            label = str(model._meta.get_field(field_name).verbose_name)
        except Exception:
            label = ""
        for error in errors:
            text = " ".join(error.messages)
            messages.append(f"{label}: {text}" if label else text)
    return messages


# ---------------------------------------------------------------- template / export

def build_workbook(products=None) -> bytes:
    """xlsx in the import format: the given products and their variants (an export to edit and re-import),
    or one sample product with sample variants."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = PRODUCTS_SHEET
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "C2"
    sheet.append([column.title for column in COLUMNS])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0D5C75")
    for index, column in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.row_dimensions[1].height = 22

    def yes_no(value):
        return "بله" if value else "خیر"

    def _jalali_or_blank(moment):
        from apps.common.admin_jalali import to_jalali

        return to_jalali(moment) if moment else ""

    if products is None:
        sheet.append(["", "یخچال فریزر نمونه مدل X", "یخچال و فریزر", "نمونه", 45000000, 42500000,
                      "", "۱۴۰۵/۰۷/۰۵ ۲۳:۵۹", 38000000, 5, "بله", "خیر",
                      "۲۴ ماه گارانتی رسمی", "ایران", "A+", 72.5, "بله", 0, 0, "", "توضیحات محصول", "fridge-1.jpg|fridge-2.jpg"])
    else:
        for p in products.select_related("category", "brand"):
            # variant products: price and stock live on the variants, so leave them out of the sheet
            priced = not p.has_variants
            sheet.append([
                p.slug, p.name, p.category.name, p.brand.name if p.brand_id else "",
                p.price if priced else "", (p.sale_price or "") if priced else "",
                _jalali_or_blank(p.sale_starts_at), _jalali_or_blank(p.sale_ends_at), p.cost_price or "",
                p.stock if priced else "", yes_no(p.is_active), yes_no(p.is_featured), p.warranty, p.country, p.energy_class,
                float(p.weight_kg) if p.weight_kg is not None else "", yes_no(p.requires_install),
                p.installment_months, p.installment_down_payment, p.video_url, p.description, "",
            ])
    def column_letter(key):
        return get_column_letter([c.key for c in COLUMNS].index(key) + 1)

    last_row = max(sheet.max_row, 2) + 500
    for key in ("price", "sale_price", "cost_price", "installment_down_payment"):  # thousands separators
        for row in range(2, last_row + 1):
            sheet[f"{column_letter(key)}{row}"].number_format = "#,##0"

    # ---- variants: one row per colour / option
    variants_sheet = workbook.create_sheet(VARIANTS_SHEET)
    variants_sheet.sheet_view.rightToLeft = True
    variants_sheet.freeze_panes = "B2"
    variants_sheet.append([column.title for column in VARIANT_COLUMNS])
    for index, column in enumerate(VARIANT_COLUMNS, start=1):
        cell = variants_sheet.cell(row=1, column=index)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        variants_sheet.column_dimensions[get_column_letter(index)].width = column.width
    if products is None:
        sample = slugify("یخچال فریزر نمونه مدل X", allow_unicode=True)
        variants_sheet.append([sample, "سفید", "#F5F5F5", "", "", 45000000, 42500000, 3, "بله", 1, "fridge-white.jpg"])
        variants_sheet.append([sample, "استیل", "#9CA3AF", "", "", 47000000, "", 2, "بله", 2, ""])
    else:
        for variant in ProductVariant.objects.filter(product__in=products).select_related("product").order_by(
            "product__name", "sort_order", "pk"
        ):
            variants_sheet.append([
                variant.product.slug, variant.color_name, variant.color_hex, variant.option, variant.sku,
                variant.price, variant.sale_price or "", variant.stock, yes_no(variant.is_active), variant.sort_order, "",
            ])

    def variant_letter(key):
        return get_column_letter([c.key for c in VARIANT_COLUMNS].index(key) + 1)

    for key in ("price", "sale_price"):
        for row in range(2, last_row + 1):
            variants_sheet[f"{variant_letter(key)}{row}"].number_format = "#,##0"

    lists = workbook.create_sheet("فهرست‌ها")
    lists.sheet_view.rightToLeft = True
    lists.append(["دسته‌بندی‌ها", "برندها", "بله/خیر"])
    categories = list(Category.objects.order_by("name").values_list("name", flat=True))
    brands = list(Brand.objects.order_by("name").values_list("name", flat=True))
    for i in range(max(len(categories), len(brands), 2)):
        lists.append([
            categories[i] if i < len(categories) else None,
            brands[i] if i < len(brands) else None,
            ("بله", "خیر")[i] if i < 2 else None,
        ])
    for letter in "ABC":
        lists.column_dimensions[letter].width = 24

    # dropdowns that still accept new names (showErrorMessage off), so a new category can be typed
    def add_list(key, source):
        validation = DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=False)
        validation.add(f"{column_letter(key)}2:{column_letter(key)}{last_row}")
        sheet.add_data_validation(validation)

    if categories:
        add_list("category", f"='فهرست‌ها'!$A$2:$A${len(categories) + 1}")
    if brands:
        add_list("brand", f"='فهرست‌ها'!$B$2:$B${len(brands) + 1}")
    for key in BOOL_KEYS:
        add_list(key, "='فهرست‌ها'!$C$2:$C$3")
    active_list = DataValidation(type="list", formula1="='فهرست‌ها'!$C$2:$C$3", allow_blank=True, showErrorMessage=False)
    active_list.add(f"{variant_letter('is_active')}2:{variant_letter('is_active')}{last_row}")
    variants_sheet.add_data_validation(active_list)

    guide = workbook.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    guide.append(["ستون", "توضیح"])
    for cell in guide[1]:
        cell.font, cell.fill = header_font, header_fill
    for column in COLUMNS:
        guide.append([column.title, column.help])
    guide.append([])
    guide.append([f"برگهٔ «{VARIANTS_SHEET}»", "هر ردیف یک رنگ یا ظرفیت از یک محصول؛ اگر محصولی تنوع دارد، قیمت و موجودی‌اش از همین برگه می‌آید."])
    for column in VARIANT_COLUMNS:
        guide.append([column.title, column.help])
    guide.append([])
    guide.append(["نکته", "در ویرایش محصول موجود، خانه خالی مقدار فعلی را تغییر نمی‌دهد و ستونی که در فایل نباشد دست نمی‌خورد."])
    guide.append(["نکته", "قیمت‌ها به تومان هستند؛ ارقام فارسی و جداکننده هزارگان هم پذیرفته می‌شود."])
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 100

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
