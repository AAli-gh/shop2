import io
import zipfile

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils.text import slugify
from PIL import Image

from apps.catalog import importer
from apps.catalog.models import Category, Product, ProductVariant

VARIANT_HEADER = ["شناسه محصول", "رنگ", "کد رنگ", "گزینه", "کد کالا", "قیمت", "قیمت با تخفیف", "موجودی", "فعال"]


def _xlsx(products=None, variants=None, only_variants=False):
    workbook = openpyxl.Workbook()
    first = workbook.active
    if only_variants:
        first.title = "تنوع‌ها"
        for row in variants:
            first.append(row)
    else:
        first.title = "محصولات"
        for row in products or [["شناسه", "نام محصول"]]:
            first.append(row)
        if variants is not None:
            sheet = workbook.create_sheet("تنوع‌ها")
            for row in variants:
                sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile("products.xlsx", buffer.getvalue())


def _run(upload, images=None, commit=True):
    return importer.run_import(upload, upload.name, images, commit=commit)


@pytest.fixture
def fridge(db):
    category = Category.objects.create(name="یخچال", slug="fridge")
    return Product.objects.create(category=category, name="یخچال ساید", slug="side", price=1, stock=0)


def test_variants_are_created_and_set_the_product_price_and_stock(fridge):
    report = _run(_xlsx(variants=[
        VARIANT_HEADER,
        ["side", "سفید", "#FFFFFF", "", "", "50,000,000", "", "5", "بله"],
        ["side", "استیل", "", "۲۴ فوت", "ST-24", 60_000_000, 45_000_000, 2, ""],
    ]))

    assert (report.variants_created, report.errors) == (2, [])
    fridge.refresh_from_db()
    assert (fridge.price, fridge.sale_price, fridge.stock) == (60_000_000, 45_000_000, 7)
    assert ProductVariant.objects.get(sku="ST-24").label == "استیل · ۲۴ فوت"


def test_existing_variants_are_found_by_colour_and_option_or_by_sku(fridge):
    white = ProductVariant.objects.create(product=fridge, color_name="سفید", price=50_000_000, stock=5)
    steel = ProductVariant.objects.create(product=fridge, color_name="استیل", sku="ST-24", price=60_000_000, stock=2)

    report = _run(_xlsx(variants=[
        VARIANT_HEADER,
        ["side", "سفید", "", "", "", "", "", "9", ""],  # only stock: the price stays
        ["side", "", "", "", "ST-24", "61000000", "", "", ""],  # found by SKU
    ]))

    assert (report.variants_updated, report.variants_created, report.errors) == (2, 0, [])
    white.refresh_from_db()
    steel.refresh_from_db()
    assert (white.price, white.stock) == (50_000_000, 9)
    assert (steel.color_name, steel.price) == ("استیل", 61_000_000)


def test_a_new_product_and_its_variants_can_share_a_file(db):
    Category.objects.create(name="یخچال", slug="fridge")

    report = _run(_xlsx(
        products=[["شناسه", "نام محصول", "دسته‌بندی", "قیمت"], ["new-fridge", "یخچال تازه", "یخچال", "1000"]],
        variants=[VARIANT_HEADER, ["new-fridge", "سفید", "", "", "", "2000", "", "3", ""]],
    ))

    assert (report.created, report.variants_created) == (1, 1)
    product = Product.objects.get(slug="new-fridge")
    assert (product.price, product.stock) == (2000, 3)


def test_problem_rows_are_reported_on_the_variants_sheet(fridge):
    report = _run(_xlsx(variants=[
        VARIANT_HEADER,
        ["missing", "سفید", "", "", "", "1", "", "1", ""],     # row 2: unknown product
        ["side", "", "", "", "", "1", "", "1", ""],            # row 3: no colour or option
        ["side", "سفید", "", "", "", "", "", "1", ""],         # row 4: new variant without price
        ["side", "مشکی", "blue", "", "", "1", "", "1", ""],    # row 5: bad colour code
        ["side", "طلایی", "", "", "", "100", "200", "1", ""],  # row 6: sale price above price
        ["side", "نقره‌ای", "", "", "", "100", "", "1", ""],   # row 7: fine
        ["side", "نقره‌ای", "", "", "", "100", "", "1", ""],   # row 8: repeats row 7
    ]))

    messages = {row.row: row.message for row in report.errors}
    assert "missing" in messages[2]
    assert "رنگ" in messages[3] and "قیمت" in messages[4]
    assert "کد رنگ" in messages[5]
    assert "کمتر" in messages[6]
    assert "ردیف 7" in messages[8]
    assert {row.sheet for row in report.errors} == {"تنوع‌ها"}
    assert report.variants_created == 1


def test_a_file_with_only_a_variants_sheet(fridge):
    report = _run(_xlsx(variants=[VARIANT_HEADER, ["side", "سفید", "", "", "", "1000", "", "2", ""]], only_variants=True))
    assert (report.variants_created, report.errors) == (1, [])


def test_the_preview_saves_nothing(fridge):
    report = _run(_xlsx(variants=[VARIANT_HEADER, ["side", "سفید", "", "", "", "1000", "", "2", ""]]), commit=False)
    assert report.variants_created == 1
    assert not ProductVariant.objects.exists()


def test_restock_subscribers_hear_when_variants_bring_stock(fridge, monkeypatch, django_capture_on_commit_callbacks):
    notified = []
    monkeypatch.setattr("apps.orders.emails.restock_available", lambda product: notified.append(product.pk))

    with django_capture_on_commit_callbacks(execute=True):
        _run(_xlsx(variants=[VARIANT_HEADER, ["side", "سفید", "", "", "", "1000", "", "3", ""]]))

    assert notified == [fridge.pk]


def test_a_variant_photo_comes_from_the_zip(fridge, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    photo = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(photo, "PNG")
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("white.png", photo.getvalue())
    header = [*VARIANT_HEADER, "تصویر"]

    _run(_xlsx(variants=[header, ["side", "سفید", "", "", "", "1000", "", "2", "", "white.png"]]),
         images=SimpleUploadedFile("images.zip", archive.getvalue()))

    assert ProductVariant.objects.get().image.name.startswith("products/variants/white")


def test_export_lists_variants_and_imports_back(fridge):
    ProductVariant.objects.create(product=fridge, color_name="سفید", price=50_000_000, stock=5, sort_order=1)
    ProductVariant.objects.create(product=fridge, color_name="استیل", sku="ST-24", price=60_000_000, stock=2, sort_order=2)

    workbook = openpyxl.load_workbook(io.BytesIO(importer.build_workbook(Product.objects.filter(pk=fridge.pk))))
    assert workbook.sheetnames == ["محصولات", "تنوع‌ها", "فهرست‌ها", "راهنما"]
    sheet = workbook["تنوع‌ها"]
    assert [c.value for c in sheet[1]] == [c.title for c in importer.VARIANT_COLUMNS]
    assert [sheet.cell(row=r, column=2).value for r in (2, 3)] == ["سفید", "استیل"]
    stock_column = [c.key for c in importer.VARIANT_COLUMNS].index("stock") + 1
    sheet.cell(row=2, column=stock_column, value=8)
    edited = io.BytesIO()
    workbook.save(edited)

    report = _run(SimpleUploadedFile("e.xlsx", edited.getvalue()))

    assert report.errors == [] and (report.updated, report.variants_updated) == (1, 2)
    fridge.refresh_from_db()
    assert fridge.stock == 10


def test_the_template_has_sample_variants_of_the_sample_product(db):
    workbook = openpyxl.load_workbook(io.BytesIO(importer.build_workbook()))
    sample_slug = slugify("یخچال فریزر نمونه مدل X", allow_unicode=True)
    assert [workbook["تنوع‌ها"].cell(row=r, column=1).value for r in (2, 3)] == [sample_slug, sample_slug]
